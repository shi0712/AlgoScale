#!/usr/bin/env python3
"""Build the public AlgoScale problem catalog from the maintained metadata files."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import openpyxl


PLATFORM_NAMES = {
    "洛谷": "Luogu",
    "qoj": "QOJ",
    "QOJ": "QOJ",
    "Codeforces": "Codeforces",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--origins-csv", type=Path, required=True)
    parser.add_argument("--complexity-xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def read_complexities(path: Path) -> dict[str, dict[str, str | None]]:
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    records: dict[str, dict[str, str | None]] = {}

    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)]
        record = dict(zip(headers, values))
        title = str(record.get("local_title") or "").strip()
        if not title:
            continue
        if title in records:
            raise ValueError(f"duplicate local title in complexity sheet: {title}")
        records[title] = {
            "originalComplexity": record.get("intended time complexity"),
            "algoscaleComplexity": record.get("advanced time complexity"),
        }
    return records


def main() -> None:
    args = parse_args()
    complexities = read_complexities(args.complexity_xlsx)
    existing_code: dict[int, dict[str, str]] = {}
    if args.output.is_file():
        existing_catalog = json.loads(args.output.read_text(encoding="utf-8"))
        existing_code = {
            int(item["id"]): item.get("generatedCode", {})
            for item in existing_catalog
            if isinstance(item, dict) and "id" in item
        }

    with args.origins_csv.open(encoding="utf-8-sig", newline="") as handle:
        origins = list(csv.DictReader(handle))

    catalog = []
    for origin in origins:
        title = origin["local_title"].strip()
        complexity = complexities.get(title, {})
        problem_id = int(origin["local_problem_id"])
        original_url = origin["original_url"].strip()
        platform = PLATFORM_NAMES.get(
            origin["source_platform"].strip(), origin["source_platform"].strip()
        )
        catalog.append(
            {
                "id": problem_id,
                "title": title,
                "originalTitle": origin["original_title"].strip(),
                "originalUrl": original_url or None,
                "platform": platform or "Pending",
                "originalComplexity": complexity.get("originalComplexity") or None,
                "algoscaleComplexity": complexity.get("algoscaleComplexity") or None,
                "track": "Interactive" if problem_id >= 82 else "Standard",
                "generatedCode": existing_code.get(problem_id, {}),
                "metadataStatus": (
                    "verified"
                    if original_url
                    and complexity.get("originalComplexity")
                    and complexity.get("algoscaleComplexity")
                    else "pending"
                ),
            }
        )

    expected_ids = list(range(1, 89))
    actual_ids = [item["id"] for item in catalog]
    if actual_ids != expected_ids:
        raise ValueError(f"expected problem IDs 1..88, got {actual_ids}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    pending = sum(item["metadataStatus"] == "pending" for item in catalog)
    print(f"wrote {len(catalog)} problems to {args.output} ({pending} pending metadata)")


if __name__ == "__main__":
    main()
